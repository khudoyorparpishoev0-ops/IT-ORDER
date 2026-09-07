"""Выгрузки в Excel и PDF."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

from app.core.text import count_with_word, days, plural
from app.db.models import PaymentMethod
from app.schemas.request import DecisionIn, ExpenseLineIn, PaymentIn, RequestCreate
from app.services import requests as svc
from app.services.export_pdf import register_fonts

XLSX_MAGIC = b"PK\x03\x04"
PDF_MAGIC = b"%PDF"


def make_paid(session, employee, manager, project, amount: str, document: str):
    request = svc.create_request(
        session,
        RequestCreate(
            employee_id=employee.id,
            project_id=project.id,
            lines=[ExpenseLineIn(title="Расход", quantity=1, price=amount)],
            submit=True,
        ),
    )
    if request.status.value == "pending":
        svc.decide_request(
            session, request.id, DecisionIn(approve=True, actor=manager.full_name)
        )
    svc.pay_request(
        session,
        request.id,
        PaymentIn(method=PaymentMethod.CARD, document=document, actor="ФИНАНСЫ"),
    )
    session.flush()
    return request


# --------------------------------------------------------------------------
# Склонение числительных
# --------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("number", "expected"),
    [
        (1, "выплата"),
        (2, "выплаты"),
        (4, "выплаты"),
        (5, "выплат"),
        (0, "выплат"),
        (11, "выплат"),
        (12, "выплат"),
        (14, "выплат"),
        (21, "выплата"),
        (22, "выплаты"),
        (25, "выплат"),
        (101, "выплата"),
        (111, "выплат"),
    ],
)
def test_plural_forms(number: int, expected: str) -> None:
    assert plural(number, "выплата", "выплаты", "выплат") == expected


def test_count_with_word() -> None:
    assert count_with_word(2, "заявка", "заявки", "заявок") == "2 заявки"
    assert days(1) == "1 день"
    assert days(5) == "5 дней"


# --------------------------------------------------------------------------
# Реестр выплат
# --------------------------------------------------------------------------
def test_payments_xlsx(client, login, employee, manager, finance, project, session) -> None:
    make_paid(session, employee, manager, project, "3100.00", "ПП-0412")
    make_paid(session, employee, manager, project, "1480.00", "ПП-0409")

    login(finance)
    response = client.get("/api/exports/payments.xlsx")
    assert response.status_code == 200
    assert response.content.startswith(XLSX_MAGIC)

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Реестр выплат"

    values = [
        [cell.value for cell in row] for row in ws.iter_rows(min_row=4, max_row=7)
    ]
    header = values[0]
    assert header[0] == "Оплачена"
    assert header[4] == "Сумма, TJS"

    # Суммы — числа, а не строки: иначе получатель не просуммирует столбец
    amounts = [row[4] for row in values[1:3]]
    assert all(isinstance(a, (int, float, Decimal)) for a in amounts), amounts
    assert sum(Decimal(str(a)) for a in amounts) == Decimal("4580.00")

    total_row = values[3]
    assert total_row[0] == "Итого выплачено"
    assert Decimal(str(total_row[4])) == Decimal("4580.00")


def test_payments_xlsx_filename_is_cyrillic(client, login, finance) -> None:
    login(finance)
    disposition = client.get("/api/exports/payments.xlsx").headers[
        "content-disposition"
    ]
    # Кириллица передаётся только через filename* по RFC 5987
    assert "filename*=UTF-8''" in disposition
    encoded = disposition.split("filename*=UTF-8''")[1]
    assert "Реестр-выплат" in unquote(encoded)
    assert unquote(encoded).endswith(".xlsx")


def test_payments_pdf(client, login, employee, manager, finance, project, session) -> None:
    make_paid(session, employee, manager, project, "3100.00", "ПП-0412")

    login(finance)
    response = client.get("/api/exports/payments.pdf")
    assert response.status_code == 200
    assert response.content.startswith(PDF_MAGIC)
    assert len(response.content) > 5000, "PDF со встроенными шрифтами не бывает крошечным"


def test_empty_register_still_exports(client, login, finance) -> None:
    """Пустой период не должен ронять выгрузку."""
    login(finance)
    assert client.get("/api/exports/payments.xlsx").status_code == 200
    assert client.get("/api/exports/payments.pdf").status_code == 200


# --------------------------------------------------------------------------
# Заявки
# --------------------------------------------------------------------------
def test_requests_xlsx(as_manager, manager, project) -> None:
    as_manager.post(
        "/api/requests",
        json={
            "employee_id": manager.id,
            "project_id": project.id,
            "lines": [{"title": "Материалы", "quantity": 2, "price": "700.00"}],
        },
    )
    response = as_manager.get("/api/exports/requests.xlsx")
    assert response.status_code == 200

    ws = load_workbook(BytesIO(response.content)).active
    assert ws.title == "Заявки"
    row = [cell.value for cell in ws[5]]
    assert row[2] == manager.full_name
    assert Decimal(str(row[5])) == Decimal("1400.00")
    assert row[6] == "На утверждении"


def test_employee_exports_only_own_requests(
    client, login, employee, manager, project
) -> None:
    """Выгрузка не должна обходить ограничение видимости."""
    login(manager)
    client.post(
        "/api/requests",
        json={
            "employee_id": manager.id,
            "project_id": project.id,
            "lines": [{"title": "Чужая", "quantity": 1, "price": "900.00"}],
        },
    )
    login(employee)
    client.post(
        "/api/requests",
        json={
            "employee_id": employee.id,
            "project_id": project.id,
            "lines": [{"title": "Своя", "quantity": 1, "price": "800.00"}],
        },
    )

    ws = load_workbook(BytesIO(client.get("/api/exports/requests.xlsx").content)).active
    names = {ws.cell(row=r, column=3).value for r in range(5, ws.max_row)}
    assert names == {employee.full_name}


def test_request_pdf(as_manager, manager, project) -> None:
    created = as_manager.post(
        "/api/requests",
        json={
            "employee_id": manager.id,
            "project_id": project.id,
            "lines": [{"title": "Такси до объекта", "quantity": 2, "price": "535.00"}],
        },
    ).json()

    response = as_manager.get(f"/api/exports/requests/{created['id']}.pdf")
    assert response.status_code == 200
    assert response.content.startswith(PDF_MAGIC)
    assert created["number"] in unquote(response.headers["content-disposition"])


def test_foreign_request_pdf_is_404(client, login, employee, manager, project) -> None:
    login(manager)
    foreign = client.post(
        "/api/requests",
        json={
            "employee_id": manager.id,
            "project_id": project.id,
            "lines": [{"title": "Чужая", "quantity": 1, "price": "900.00"}],
        },
    ).json()
    login(employee)
    assert client.get(f"/api/exports/requests/{foreign['id']}.pdf").status_code == 404


# --------------------------------------------------------------------------
# Права
# --------------------------------------------------------------------------
def test_employee_cannot_export_payments(client, login, employee) -> None:
    login(employee)
    assert client.get("/api/exports/payments.xlsx").status_code == 403
    assert client.get("/api/exports/payments.pdf").status_code == 403


def test_anonymous_cannot_export(client) -> None:
    for path in (
        "/api/exports/payments.xlsx",
        "/api/exports/payments.pdf",
        "/api/exports/requests.xlsx",
    ):
        assert client.get(path).status_code == 401, path


# --------------------------------------------------------------------------
# Шрифты
# --------------------------------------------------------------------------
def test_pdf_fonts_cover_russian_and_tajik() -> None:
    """Пропущенный глиф печатается чёрным прямоугольником — это брак."""
    from reportlab.pdfbase import pdfmetrics

    register_fonts()
    probe = "АБВЯабвя ғӣқӯҳҷ ABCabc 0123 №«»—"
    for name in ("Manrope", "Manrope-Bold", "JetBrainsMono", "JetBrainsMono-SemiBold"):
        face = pdfmetrics.getFont(name).face
        missing = [c for c in probe if c != " " and ord(c) not in face.charToGlyph]
        assert not missing, f"{name}: нет глифов {''.join(missing)}"
